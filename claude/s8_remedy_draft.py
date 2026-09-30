# S8: 조치단계 초안 재작성 — fault-cases.json(현재) + 측정지점(MP_*) 연결 반영. 결과: claude/remedySteps_draft.json, 고장사례_조치단계_초안.xlsx (둘 다 untracked)
import json, re
import openpyxl
fc = json.load(open('fault-cases.json', encoding='utf-8'))['cases']
M = {p['id'] for p in json.load(open('measure-points.json', encoding='utf-8'))['points']}
B = {x['id'] for l in json.load(open('board-info.json', encoding='utf-8'))['boards'].values() for x in l}
P = lambda *a: ['MP_' + x for x in a]
# 여러 단계 사례의 단계별 측정지점(단계 번호 → MP). 없는 단계는 측정지점 없음. 단일 단계 사례는 사례의 MP 전부.
STEP_MP = {
    'FLT_ECD_02': {1: P('ECD_STEKOP_LED', 'ECD_STEKOP_24V'), 2: P('ECD_STEKOP_LED')},
    'FLT_ECD_03': {2: P('ECD_DEWEMO_LED')},
    'FLT_ECD_08': {2: P('ECD_DEWEMO_MOTOR110', 'FIELD_SW_MOTOR')}, 'FLT_ECD_09': {2: P('ECD_DEWEMO_MOTOR110', 'FIELD_SW_MOTOR')},
    'FLT_ILK_07': {1: P('SICAS_S51_DC60')}, 'FLT_PWR_04': {1: P('SICAS_S51_DC60')},
    'FLT_PWR_06': {1: P('ECD_SV2602_OUT')},
    'FLT_SIG_06': {1: P('ECD_STEKOP_LED', 'ECD_STEKOP_24V'), 2: P('ECD_STEKOP_LED')},
    'FLT_ATO_01': {3: P('ATO_LOOP_CURRENT'), 4: P('ATO_LOOP_CURRENT')},
    'FLT_ATO_02': {2: P('ATO_LOOP_CURRENT')}, 'FLT_ATO_03': {2: P('ATO_LOOP_CURRENT')},
    'FLT_ATO_04': {1: P('ATO_CONTROLLER_LED_TEST'), 4: P('ATO_CONTROLLER_LED_TEST')},
    'FLT_VEN_01': {1: P('SICAS_VENUS3_LED')},
    'FLT_TWC_01': {1: P('TWC_OUT')},
    'FLT_TWC_02': {1: P('TWC_RX_WAVEFORM'), 2: P('TWC_IMU100_LED')},
    'FLT_TWC_03': {1: P('TWC_CB_DETAIL'), 2: P('TWC_IMU100_LED')},
    'FLT_FTGS_04': {1: P('FTGS_TRACK_CB'), 2: P('FTGS_TRACK_CB'), 3: P('FTGS_DCBOX', 'FTGS_TX_FIELD', 'FTGS_RX_FIELD')},
}
BOARD_KW = [(r'STEKOP', 'ECDSTT_STEKOP'), (r'DEWEMO', 'ECDSTT_DEWEMO'), (r'DESIMO', 'ECDSTT_DESIMO'), (r'SV ?2602', 'ECDSTT_PS'),
            (r'VENUS3', 'SICAS_VENUS3'), (r'SVK ?2102', 'SICAS_SVK2102'), (r'FUUELL', 'ATOLOOP_FUUELL'), (r'REMEMO', 'ATOLOOP_REMEMO'),
            (r'TWC 서브랙|모듈을 예비품|IMU', 'TWC_IMU100'), (r'GF계전기', 'FTGS_GF_RELAY_PCB'), (r'파워 서플라이|전원장치', None)]
CASE_BOARD = {'FLT_FTGS_02': 'FTGS_POWER_UNIT', 'FLT_ILK_01': 'SICAS_OLM', 'FLT_ILK_03': 'SICAS_OLM', 'FLT_SIG_02': 'SICAS_OLM'}


def board(cid, text):
    for pat, b in BOARD_KW:
        if re.search(pat, text) and b: return b
    if cid in CASE_BOARD and re.search(r'플러그|라인|전원장치|파워', text): return CASE_BOARD[cid]
    return None


out = {}
for c in fc:
    cid = c['id']; mps = c.get('measurePointRef') or []; rem = c['remedy']; steps = []
    for i, a in enumerate(rem, 1):
        if cid in STEP_MP: m = STEP_MP[cid].get(i, [])
        else: m = list(mps) if len(rem) == 1 else []
        assert all(x in M for x in m), (cid, m)
        b = board(cid, a); assert b is None or b in B, (cid, b)
        steps.append({'order': i, 'action': a, 'measurePointRef': m or None, 'boardRef': b})
    if any(s['measurePointRef'] or s['boardRef'] for s in steps): out[cid] = steps
json.dump(out, open('claude/remedySteps_draft.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
F = {c['id']: c for c in fc}
f1 = '고장사례_조치단계_초안.xlsx'; wb = openpyxl.load_workbook(f1); ws = wb['조치단계_초안']
for i in range(ws.max_row, 1, -1): ws.delete_rows(i)
for cid, st in out.items():
    for s in st:
        ws.append([cid, F[cid]['symptom'], F[cid]['domain'], s['order'], s['action'], ', '.join(s['measurePointRef'] or []) or None, s['boardRef'], None, None])
wb.save(f1)
print('cases', len(out), 'steps', ws.max_row - 1, 'with MP', sum(1 for st in out.values() for s in st if s['measurePointRef']), 'with board', sum(1 for st in out.values() for s in st if s['boardRef']))
